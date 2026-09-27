"""Excel for Power BI: 3 sheets (Tanker, Oil, Oil Price) joined on Date (MM/DD/YYYY).

    python3 scripts/build_excel.py                       # 09/20/2026 → 09/27/2026
    python3 scripts/build_excel.py --start 2026-09-20 --end 2026-09-27

Sheets (every sheet has a real Excel date column "Date", format mm/dd/yyyy):
  Tanker     one row per vessel seen by VesselAPI (western Musandam coast) in the period.
             Date = first day it was seen; Latitude/Longitude = its last REAL position that day.
             Vessel type from VesselAPI's vessel record; Unknown if we have none (never guessed).
  Oil        one row per day. Cargo per ship is NOT available (AIS has no cargo), so this is per day:
             IMF PortWatch tanker transits + tanker trade-volume estimate for the whole strait,
             converted to barrels with 7.33 bbl/t (crude average) — an ESTIMATE.
             Days PortWatch has not published yet stay empty.
  Oil Price  one row per day. Brent spot price (EIA via FRED, daily). This is NOT an opening price —
             no free official source for the opening price was found. Weekends: no price.

Everything is real data; empty cells mean "not available", never a made-up value.
Output: samples/powerbi/hormuz_watch_<start>_<end>.xlsx (git-ignored: contains VesselAPI data).
"""
import argparse
import csv
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from export_daily_vessels import crossings_by_day  # noqa: E402
from export_powerbi import TEMP_GATE, is_tanker, load_details, location_text, tanker_class  # noqa: E402
from src.history import day_coverage, normalize_all, read_sources  # noqa: E402

BARRELS_PER_TONNE = 7.33          # average crude oil conversion (estimate)
BRENT_CACHE = ROOT / "samples" / "brent_fred.csv"
DATE_FMT = "mm/dd/yyyy"


# ---------- data loading (network parts are optional; failures leave cells empty) ----------

def load_portwatch_rows(refresh: bool) -> tuple:
    import fetch_portwatch as pw
    note = ""
    if refresh or not pw.OUT.exists():
        try:
            rows = pw.fetch(60)
            pw.OUT.parent.mkdir(exist_ok=True)
            with open(pw.OUT, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=["date"] + pw.FIELDS, extrasaction="ignore")
                w.writeheader()
                w.writerows(rows)
        except Exception as e:  # network / service problem: keep going, cells stay empty
            note = f"PortWatch không tải được: {e}"
    return pw.load(), note


def load_brent(refresh: bool) -> tuple:
    from src.brent_loader import fetch_fred as fetch_brent
    note = ""
    if refresh or not BRENT_CACHE.exists():
        try:
            rows = fetch_brent(120)
            BRENT_CACHE.parent.mkdir(exist_ok=True)
            with open(BRENT_CACHE, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=["price_date", "price_usd"])
                w.writeheader()
                w.writerows(rows)
        except Exception as e:
            note = f"FRED/EIA không tải được: {e}"
    prices = {}
    if BRENT_CACHE.exists():
        with open(BRENT_CACHE, encoding="utf-8-sig") as f:
            prices = {r["price_date"]: float(r["price_usd"]) for r in csv.DictReader(f)}
    return prices, note


def _num(value):
    try:
        return float(value) if value not in (None, "") else None
    except ValueError:
        return None


# ---------- sheet rows ----------

def tanker_rows(positions: list, details: dict, crossings: dict, start: date, end: date) -> list:
    by_vessel = {}
    for p in positions:
        d = date.fromisoformat(p["timestamp_utc"][:10])
        if start <= d <= end:
            by_vessel.setdefault(p["mmsi"], []).append(p)
    rows = []
    for mmsi, track in by_vessel.items():
        first_day = track[0]["timestamp_utc"][:10]
        same_day = [p for p in track if p["timestamp_utc"][:10] == first_day]
        last = same_day[-1]
        v = details.get(mmsi) or {}
        cls = tanker_class(details.get(mmsi))
        days = sorted({p["timestamp_utc"][:10] for p in track})
        name = next((p["ship_name"] for p in reversed(track) if p.get("ship_name")), None) or v.get("name")
        rows.append({
            "Date": date.fromisoformat(first_day),
            "MMSI": mmsi,
            "Ship name": name or "",
            "Registry name": v.get("name") or "",
            "Vessel type": v.get("vessel_type") or "Unknown",
            "Tanker class": cls,
            "Is tanker": is_tanker(cls),
            "Flag": v.get("country") or "",
            "Deadweight (t)": v.get("deadweight_tonnage") if v.get("deadweight_tonnage") else None,
            "Latitude": round(last["latitude"], 5),
            "Longitude": round(last["longitude"], 5),
            "Location": location_text(last["latitude"], last["longitude"]),
            "Time seen (UTC)": last["timestamp_utc"][11:16],
            "Days observed": ", ".join(datetime.strptime(d, "%Y-%m-%d").strftime("%m/%d/%Y") for d in days),
            "Crossed temp gate": "; ".join(c for (d, m), c in sorted(crossings.items()) if m == mmsi and d in days),
            "Data": "REAL (VesselAPI)",
        })
    rows.sort(key=lambda r: (r["Date"], r["Is tanker"] != "Yes", r["Ship name"]))
    return rows


def oil_rows(portwatch: dict, positions: list, details: dict, intervals: list, start: date, end: date,
             downloaded: bool = True) -> list:
    latest = max(portwatch) if portwatch else None
    rows, day = [], start
    while day <= end:
        pw = portwatch.get(day.isoformat())
        tonnes = _num(pw.get("capacity_tanker")) if pw else None
        seen = {p["mmsi"] for p in positions if p["timestamp_utc"][:10] == day.isoformat()}
        classes = [tanker_class(details.get(m)) for m in seen]
        observed = sum((e - s).total_seconds() for s, e, _ in day_coverage(intervals, day)) / 3600
        rows.append({
            "Date": day,
            "Tanker transits (PortWatch)": int(_num(pw["n_tanker"])) if pw and _num(pw.get("n_tanker")) is not None else None,
            "All transits (PortWatch)": int(_num(pw["n_total"])) if pw and _num(pw.get("n_total")) is not None else None,
            "Tanker volume est. (t, PortWatch)": tonnes,
            "Oil volume est. (barrels)": round(tonnes * BARRELS_PER_TONNE) if tonnes is not None else None,
            "Tankers seen (VesselAPI)": sum(c != "Not tanker" and c != "Unknown" for c in classes),
            "Oil tankers seen (VesselAPI)": sum(c == "Oil tanker" for c in classes),
            "VesselAPI observed hours (UTC)": round(observed, 1),
            "Status": ("PortWatch published" if pw else
                       "PortWatch not downloaded" if not downloaded else
                       "PortWatch not published yet" if latest is None or day.isoformat() > latest else
                       "PortWatch no data"),
            "Source": "IMF PortWatch (satellite AIS, whole strait); VesselAPI (west Musandam coast)",
        })
        day += timedelta(days=1)
    return rows


def price_rows(prices: dict, start: date, end: date, downloaded: bool = True) -> list:
    latest = max(prices) if prices else None
    rows, day = [], start
    while day <= end:
        price = prices.get(day.isoformat())
        weekend = day.weekday() >= 5
        status = ("Price available" if price is not None else
                  "Weekend - market closed" if weekend else
                  "Not downloaded" if not downloaded else
                  "Not published yet" if latest is None or day.isoformat() > latest else
                  "No price (holiday)")
        rows.append({"Date": day, "Brent (USD/bbl)": price, "Price type": "Brent spot, daily (EIA) - not opening price",
                     "Status": status, "Source": "EIA via FRED (DCOILBRENTEU)"})
        day += timedelta(days=1)
    return rows


# ---------- Excel ----------

def write_xlsx(path: Path, sheets: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        cols = list(rows[0].keys())
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r[c] for c in cols])
        for idx, col in enumerate(cols, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            if col == "Date":
                for cell in ws[letter][1:]:
                    cell.number_format = DATE_FMT
            width = max([len(col)] + [len(str(r[col])) for r in rows if r[col] is not None])
            ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    p = argparse.ArgumentParser(description="Excel 3 sheet (Tanker, Oil, Oil Price) cho Power BI.")
    p.add_argument("--start", default="2026-09-20")
    p.add_argument("--end", default="2026-09-27")
    p.add_argument("--refresh", action="store_true", help="tải lại PortWatch và giá Brent")
    args = p.parse_args()
    start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)

    records, intervals, _ = read_sources()
    positions, _ = normalize_all(records)
    details = load_details(ROOT / "samples" / "vessel_details.jsonl")
    portwatch, pw_note = load_portwatch_rows(args.refresh)
    prices, price_note = load_brent(args.refresh)

    sheets = {"Tanker": tanker_rows(positions, details, crossings_by_day(positions, TEMP_GATE), start, end),
              "Oil": oil_rows(portwatch, positions, details, intervals, start, end, not pw_note),
              "Oil Price": price_rows(prices, start, end, not price_note)}
    out = ROOT / "samples" / "powerbi" / f"hormuz_watch_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    write_xlsx(out, sheets)

    t = sheets["Tanker"]
    print(f"Tanker: {len(t)} tàu ({sum(r['Is tanker'] == 'Yes' for r in t)} tanker, "
          f"{sum(r['Is tanker'] == 'Unknown' for r in t)} chưa rõ loại)")
    print(f"Oil: {sum(r['Status'] == 'PortWatch published' for r in sheets['Oil'])}/{len(sheets['Oil'])} ngày có số PortWatch")
    print(f"Oil Price: {sum(r['Brent (USD/bbl)'] is not None for r in sheets['Oil Price'])}/{len(sheets['Oil Price'])} ngày có giá")
    for note in (pw_note, price_note):
        if note:
            print("CẢNH BÁO:", note)
    print(f"Đã ghi {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
