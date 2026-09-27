"""Tests for scripts/build_excel.py (no network). Records and prices are SIMULATED examples."""
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import build_excel as bx  # noqa: E402

POS = [  # one simulated vessel seen on two days
    {"mmsi": "999000001", "ship_name": "SIM-1", "latitude": 26.27708, "longitude": 56.18, "timestamp_utc": "2026-09-20T04:49:00Z"},
    {"mmsi": "999000001", "ship_name": "SIM-1", "latitude": 26.26903, "longitude": 56.18469, "timestamp_utc": "2026-09-20T05:00:00Z"},
    {"mmsi": "999000001", "ship_name": "SIM-1", "latitude": 26.3, "longitude": 56.2, "timestamp_utc": "2026-09-27T01:00:00Z"},
]
DETAILS = {"999000001": {"name": "SIM-1", "vessel_type": "Crude Oil Tanker", "deadweight_tonnage": 100000}}


def test_one_row_per_vessel_with_last_position_of_first_day():
    rows = bx.tanker_rows(POS, DETAILS, {}, date(2026, 9, 20), date(2026, 9, 27))
    assert len(rows) == 1
    r = rows[0]
    assert r["Date"] == date(2026, 9, 20) and r["Location"] == "26.2690°N, 56.1847°E"
    assert (r["Latitude"], r["Longitude"], r["Time seen (UTC)"]) == (26.26903, 56.18469, "05:00")
    assert r["Days observed"] == "09/20/2026, 09/27/2026" and r["Tanker class"] == "Oil tanker"


def test_unknown_type_is_not_guessed():
    r = bx.tanker_rows(POS, {}, {}, date(2026, 9, 20), date(2026, 9, 27))[0]
    assert (r["Vessel type"], r["Is tanker"], r["Deadweight (t)"]) == ("Unknown", "Unknown", None)


def test_oil_per_day_uses_portwatch_and_marks_unpublished_days():
    pw = {"2026-09-20": {"n_tanker": "2", "n_total": "7", "capacity_tanker": "100000"}}
    rows = bx.oil_rows(pw, POS, DETAILS, [], date(2026, 9, 20), date(2026, 9, 22))
    assert [r["Date"] for r in rows] == [date(2026, 9, 20), date(2026, 9, 21), date(2026, 9, 22)]
    assert rows[0]["Tanker transits (PortWatch)"] == 2 and rows[0]["Oil volume est. (barrels)"] == 733000
    assert rows[0]["Oil tankers seen (VesselAPI)"] == 1
    assert rows[1]["Oil volume est. (barrels)"] is None and rows[1]["Status"] == "PortWatch not published yet"
    failed = bx.oil_rows({}, POS, DETAILS, [], date(2026, 9, 20), date(2026, 9, 20), downloaded=False)
    assert failed[0]["Status"] == "PortWatch not downloaded"


def test_price_rows_weekend_and_missing_days():
    prices = {"2026-09-21": 113.5, "2026-09-22": 114.89}
    rows = {r["Date"]: r for r in bx.price_rows(prices, date(2026, 9, 20), date(2026, 9, 23))}
    assert rows[date(2026, 9, 20)]["Status"] == "Weekend - market closed"      # Sunday
    assert rows[date(2026, 9, 22)]["Brent (USD/bbl)"] == 114.89
    assert rows[date(2026, 9, 23)]["Status"] == "Not published yet"


def test_excel_dates_are_real_dates_with_same_format(tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "x.xlsx"
    bx.write_xlsx(out, {"Tanker": bx.tanker_rows(POS, DETAILS, {}, date(2026, 9, 20), date(2026, 9, 27)),
                        "Oil Price": bx.price_rows({}, date(2026, 9, 20), date(2026, 9, 21))})
    wb = load_workbook(out)
    for ws in wb.worksheets:
        assert ws["A1"].value == "Date" and ws["A2"].number_format == "mm/dd/yyyy"
        assert ws["A2"].is_date
